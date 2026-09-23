"""A12.2 step 1 — parse an uploaded vendor price-list file (.xlsx or .csv)
into a header row plus a list of raw row dicts, keyed by the file's own
column headers (not yet mapped to canonical fields). `.xls` is refused.

Header detection: the first row (scanning at most 10) whose non-empty cell
count is >= 4. The PH Store template carries a second header row of canonical
identifiers (`vendor_sku`, `price_to_pharmacy`, ...) immediately under the
human-readable Portuguese labels — recognised and skipped automatically so
data starts on the row after it.
"""
import csv
import hashlib
import io
from pathlib import Path

import openpyxl

from rova.core.errors import ApiError
from rova.core.storage import save_bytes

MAX_ROWS = 5000  # [assume: A-BE-05]
MAX_BYTES = 5 * 1024 * 1024  # [assume: A-BE-05]

# canonical field identifiers as they appear in the template's Dicionário /
# second header row — used to recognise and skip that row automatically.
_CANONICAL_FIELDS = frozenset({
    "vendor_sku", "product_name_brand", "product_name_inn", "form", "strength", "pack_size",
    "pack_unit", "manufacturer", "country_of_origin", "aim_number", "barcode", "price_to_pharmacy",
    "price_to_public", "discount_percent", "cost_to_us", "stock_status", "available_quantity",
    "batch_number", "expiry_date", "min_order_qty", "currency", "effective_from", "notes",
})


class ParsedFile:
    def __init__(self, headers: list[str], rows: list[dict], header_row_index: int, stored_ref: str, ext: str):
        self.headers = headers
        self.rows = rows  # list of {header: raw cell value}
        self.header_row_index = header_row_index
        self.stored_ref = stored_ref
        self.ext = ext


def _non_empty_count(row: list) -> int:
    return sum(1 for c in row if c not in (None, ""))


def _detect_header_row(grid: list[list]) -> int:
    for i, row in enumerate(grid[:10]):
        if _non_empty_count(row) >= 4:
            return i
    raise ApiError("VALIDATION_ERROR", "could not detect a header row in the first 10 rows of the file")


def _is_canonical_identifier_row(row: list[str]) -> bool:
    non_empty = [c for c in row if c not in (None, "")]
    if not non_empty:
        return False
    hits = sum(1 for c in non_empty if str(c).strip().lower() in _CANONICAL_FIELDS)
    return hits >= max(1, len(non_empty) // 2)


def _parse_xlsx(data: bytes) -> tuple[list[str], list[dict], int]:
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    sheet_name = "Lista de Preços" if "Lista de Preços" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]
    grid = [[c.value for c in row] for row in ws.iter_rows(max_row=min(ws.max_row, MAX_ROWS + 20))]
    if not grid:
        raise ApiError("VALIDATION_ERROR", "the file has no rows")
    header_idx = _detect_header_row(grid)
    headers = [str(h).strip() if h is not None else "" for h in grid[header_idx]]
    data_start = header_idx + 1
    if data_start < len(grid) and _is_canonical_identifier_row([str(c) if c is not None else "" for c in grid[data_start]]):
        data_start += 1
    rows = []
    for raw_row in grid[data_start:]:
        if _non_empty_count(raw_row) == 0:
            continue
        row = {headers[i]: (raw_row[i] if i < len(raw_row) else None) for i in range(len(headers)) if headers[i]}
        rows.append(row)
    return headers, rows, header_idx


def _sniff_csv_dialect(text: str) -> str:
    sample = text[:4096]
    try:
        return csv.Sniffer().sniff(sample, delimiters=";,").delimiter
    except csv.Error:
        return ";" if sample.count(";") > sample.count(",") else ","


def _parse_csv(data: bytes) -> tuple[list[str], list[dict], int]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    delim = _sniff_csv_dialect(text)
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    grid = list(reader)
    if not grid:
        raise ApiError("VALIDATION_ERROR", "the file has no rows")
    header_idx = _detect_header_row(grid)
    headers = [h.strip() for h in grid[header_idx]]
    data_start = header_idx + 1
    if data_start < len(grid) and _is_canonical_identifier_row(grid[data_start]):
        data_start += 1
    rows = []
    for raw_row in grid[data_start:]:
        if _non_empty_count(raw_row) == 0:
            continue
        row = {headers[i]: (raw_row[i] if i < len(raw_row) else None) for i in range(len(headers)) if headers[i]}
        rows.append(row)
    return headers, rows, header_idx


def parse(filename: str, data: bytes, *, vendor_id: str) -> ParsedFile:
    if len(data) > MAX_BYTES:
        raise ApiError("VALIDATION_ERROR", f"file exceeds the {MAX_BYTES // (1024*1024)} MB limit")
    ext = Path(filename).suffix.lower()
    if ext == ".xls":
        raise ApiError("VALIDATION_ERROR", ".xls is not accepted; please re-save as .xlsx or .csv")
    if ext == ".xlsx":
        headers, rows, header_idx = _parse_xlsx(data)
    elif ext == ".csv":
        headers, rows, header_idx = _parse_csv(data)
    else:
        raise ApiError("VALIDATION_ERROR", "only .xlsx or .csv files are accepted")
    if len(rows) > MAX_ROWS:
        raise ApiError("VALIDATION_ERROR", f"file has {len(rows)} rows; the limit is {MAX_ROWS}")
    stored_ref = save_bytes(f"pricelists/{vendor_id}", filename, data)
    return ParsedFile(headers=headers, rows=rows, header_row_index=header_idx, stored_ref=stored_ref, ext=ext)


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()
