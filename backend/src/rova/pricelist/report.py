"""A12.2 step 4 (`report.py`) — one xlsx sheet, columns `linha, texto
original, resultado, motivo (pt)` (§3 of مواصفة_استيراد_قائمة_الأسعار.md: the
line number is the VENDOR's own row number, never an internal id), so the
error report the vendor sees names exactly the row he typed."""
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import text
from sqlalchemy.orm import Session

from rova.core.storage import save_bytes

_HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_FILLS = {
    "ACCEPTED": PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
    "WARNING": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
    "PENDING_REVIEW": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
    "REJECTED": PatternFill(start_color="FCE4E4", end_color="FCE4E4", fill_type="solid"),
}


def build_report(session: Session, version_id: str, *, vendor_id: str) -> str:
    rows = session.execute(
        text(
            "SELECT row_index, raw_values, outcome, outcome_reason, outcome_detail FROM price_list_row "
            "WHERE version_id=:v ORDER BY row_index"
        ),
        {"v": version_id},
    ).mappings().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Relatório"
    headers = ["linha", "texto original", "resultado", "motivo (pt)"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    ws.freeze_panes = "A2"

    for i, row in enumerate(rows, start=2):
        original = ", ".join(f"{k}={v}" for k, v in (row["raw_values"] or {}).items() if v not in (None, ""))
        reason = row["outcome_reason"] or ""
        detail = row["outcome_detail"] or ""
        motivo = f"{reason}: {detail}" if reason and detail else (reason or detail or "aceite")
        values = [row["row_index"], original, row["outcome"], motivo]
        fill = _FILLS.get(row["outcome"])
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=i, column=c, value=v)
            if fill:
                cell.fill = fill
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 60

    buf = io.BytesIO()
    wb.save(buf)
    return save_bytes(f"pricelists/{vendor_id}/reports", f"{version_id}.xlsx", buf.getvalue())
